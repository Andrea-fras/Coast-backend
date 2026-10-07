#!/usr/bin/env python3
"""Real upload/outline routes; isolated storage and no provider calls."""
import json
import time
import uuid
import unittest
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
import test_http_integrity as fixture
from test_http_integrity import client, SessionLocal, ROOT, TMP, server
from database import FolderSource, SourceUpload, LearningJob, CourseOutline
from coast_content_oma.extraction import extract_pages
from coast_content_oma.normalized_source import load_pages
import fitz
import upload_lifecycle


def pdf_bytes():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((60, 60), 'Queues preserve arrival order. A stack removes the latest item first.')
        return doc.tobytes()


async def read_in_thread(pool, path):
    import asyncio
    from coast_content_oma.extraction import extract_and_cache
    return await asyncio.get_running_loop().run_in_executor(pool, extract_and_cache, path)


class UploadFlow(unittest.TestCase):
    def setUp(self):
        fixture.HttpIntegrity.setUp(self)
        self.data = pdf_bytes()
        self.url = '/api/folders/UploadCheck'
        self.oma = patch('oma_provider.is_oma_enabled', return_value=True)
        self.wake = patch('learning_jobs.wake')
        self.oma.start(); self.wake.start()
        self.addCleanup(self.oma.stop); self.addCleanup(self.wake.stop)
        # Production reads files in a separate process; here in a thread, so tests can patch extraction.
        reader = ThreadPoolExecutor(max_workers=2)
        self.addCleanup(reader.shutdown)
        self.reader = patch('server._read_upload', new=lambda path: read_in_thread(reader, path))
        self.reader.start(); self.addCleanup(self.reader.stop)

    def reserve(self, names, data=None):
        data = self.data if data is None else data
        files = [{'upload_id': uuid.uuid4().hex, 'filename': name, 'size_bytes': len(data)} for name in names]
        response = client.post(self.url + '/uploads', headers=self.h, json={'files': files})
        self.assertEqual(response.status_code, 200, response.text)
        return files

    def upload(self, file, data=None, headers=None):
        return client.post(self.url + '/upload', headers=headers or self.h,
            data={'upload_id': file['upload_id']}, files={'file': (file['filename'], self.data if data is None else data)})

    def rows(self):
        response = client.get(self.url + '/uploads', headers=self.h)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['uploads']

    def test_five_file_selection_survives_return_and_blocks_early_roadmap(self):
        files = self.reserve([f'Lecture {i}.pdf' for i in range(5)])
        ids = []
        for file in files[:4]:
            result = self.upload(file)
            self.assertEqual(result.status_code, 200, result.text)
            ids.append(result.json()['source_id'])
        self.assertEqual(len(self.rows()), 5)
        self.assertEqual([r['status'] for r in self.rows()].count('queued'), 1)
        with patch('lesson._call_llm_for_outline') as model:
            result = client.post(self.url + '/outline', headers=self.h, json={'source_ids': ids})
            self.assertEqual(result.status_code, 409, result.text)
            model.assert_not_called()
        fifth = self.upload(files[4])
        self.assertEqual(fifth.status_code, 200, fifth.text)
        self.assertTrue(all(r['status'] == 'complete' for r in self.rows()))
        with patch('lesson._call_llm_for_outline') as model:
            result = client.post(self.url + '/outline', headers=self.h, json={'source_ids': ids})
            self.assertEqual(result.status_code, 409, result.text)
            model.assert_not_called()

    def test_successful_retry_is_idempotent(self):
        file, = self.reserve(['Retry.pdf'])
        first = self.upload(file)
        self.assertEqual(first.status_code, 200, first.text)
        with patch('coast_content_oma.extraction.extract_pages', side_effect=AssertionError('Must not extract twice')):
            second = self.upload(file)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(first.json(), second.json())
        with SessionLocal() as db:
            self.assertEqual(db.query(FolderSource).count(), 1)
            self.assertEqual(db.query(LearningJob).count(), 1)

    def test_corrupt_pptx_is_retryable_and_can_be_removed(self):
        file, = self.reserve(['Corrupt.pptx'], b'invalid zip')
        failed = self.upload(file, b'invalid zip')
        self.assertEqual(failed.status_code, 400, failed.text)
        self.assertEqual(self.rows()[0]['status'], 'failed')
        blocked = client.post(self.url + '/outline', headers=self.h)
        self.assertEqual(blocked.status_code, 409, blocked.text)
        removed = client.delete(self.url + '/uploads/' + file['upload_id'], headers=self.h)
        self.assertEqual(removed.status_code, 200, removed.text)
        with SessionLocal() as db:
            upload_lifecycle.assert_ready(db, 1, 'UploadCheck')
            self.assertEqual(db.query(FolderSource).count(), 0)

    def test_ownership_and_expired_upload_retry(self):
        file, = self.reserve(['Private.pdf'])
        self.assertEqual(client.get(self.url + '/uploads', headers=self.other).json()['uploads'], [])
        self.assertEqual(client.delete(self.url + '/uploads/' + file['upload_id'], headers=self.other).status_code, 404)
        self.assertEqual(self.upload(file, headers=self.other).status_code, 409)
        with SessionLocal() as db:
            db.query(SourceUpload).update({'expires_at': time.time() - 1}); db.commit()
        self.assertEqual(self.rows()[0]['status'], 'failed')
        retry = client.post(self.url + '/uploads', headers=self.h, json={'files': [file]})
        self.assertEqual(retry.status_code, 200, retry.text)
        self.assertEqual(self.upload(file).status_code, 200)

    def test_remove_during_extraction_prevents_late_publish(self):
        file, = self.reserve(['Slow.pdf'])
        started, release = threading.Event(), threading.Event()
        def slow(*args, **kwargs):
            started.set()
            if not release.wait(8): raise TimeoutError('Test did not release extraction')
            return extract_pages(*args, **kwargs)
        before = set(server.FOLDER_UPLOADS_DIR.iterdir())
        with patch('coast_content_oma.extraction.extract_pages', side_effect=slow), ThreadPoolExecutor(max_workers=1) as pool:
            pending = pool.submit(self.upload, file)
            try:
                self.assertTrue(started.wait(4))
                duplicate = self.upload(file)
                self.assertEqual(duplicate.status_code, 409, duplicate.text)
                self.assertEqual(self.rows()[0]['status'], 'processing')
                removed = client.delete(self.url + '/uploads/' + file['upload_id'], headers=self.h)
                self.assertEqual(removed.status_code, 200, removed.text)
            finally: release.set()
            result = pending.result(timeout=8)
        self.assertEqual(result.status_code, 409, result.text)
        self.assertEqual(self.rows()[0]['status'], 'cancelled')
        with SessionLocal() as db: self.assertEqual(db.query(FolderSource).count(), 0)
        self.assertEqual(set(server.FOLDER_UPLOADS_DIR.iterdir()), before)

    def test_source_changes_during_model_call_keep_existing_roadmap(self):
        self.url = '/api/folders/Physics'
        file, = self.reserve(['A.pdf'])
        first = self.upload(file)
        self.assertEqual(first.status_code, 200, first.text)
        with SessionLocal() as db:
            original = db.query(CourseOutline).filter_by(user_id=1, folder_name='Physics').first().outline_json
        for change in ('pending', 'completed'):
            def model(*args):
                newfile, = self.reserve([change + '.pdf'])
                if change == 'completed':
                    with patch('oma_provider.is_oma_enabled', return_value=True):
                        self.assertEqual(self.upload(newfile).status_code, 200)
                return [{'title': 'Replacement', 'estimated_minutes': 20}]
            with patch('oma_provider.is_oma_enabled', return_value=False), patch('rag.build_folder_context', return_value=''), \
                 patch('lesson._call_llm_for_outline', side_effect=model):
                result = client.post(self.url + '/outline', headers=self.h)
            self.assertEqual(result.status_code, 409, result.text)
            with SessionLocal() as db:
                self.assertEqual(db.query(CourseOutline).filter_by(user_id=1, folder_name='Physics').first().outline_json, original)
                db.query(SourceUpload).filter_by(status='queued').update({'status': 'cancelled'}); db.commit()

    def test_figures_are_streamed_to_disk_never_held_together(self):
        # A lecture's figures decoded at once took 360 MB; now each goes to disk as it is read.
        import fitz
        from PIL import Image
        from coast_content_oma.extraction import extract_and_cache
        from coast_content_oma.normalized_source import LazyImage, cache_dir, load_pages
        pdf = Path(TMP.name) / "figures.pdf"
        doc = fitz.open()
        for n in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"Slide {n} about graphs")
            img = Image.new("RGB", (400, 300), (40 * n, 120, 200))
            for x in range(0, 400, 20):
                for y in range(300):
                    img.putpixel((x, y), (255, 255, 255))  # stripes: not an empty asset
            path = Path(TMP.name) / f"fig{n}.png"
            img.save(path)
            page.insert_image(fitz.Rect(72, 120, 472, 420), filename=str(path))
        doc.save(str(pdf))
        texts = extract_and_cache(str(pdf))
        self.assertEqual([t["text"].strip() for t in texts], [f"Slide {n} about graphs" for n in range(3)])
        pages = load_pages(pdf)
        figures = [im["pil_image"] for page in pages for im in page["images"]]
        self.assertEqual(len(figures), 3)
        self.assertTrue(all(isinstance(f, LazyImage) for f in figures))
        self.assertEqual(figures[0].size, (400, 300))
        self.assertEqual(figures[0].load().size, (400, 300))  # pixels on demand, owned by the caller
        named = {im["file"] for row in json.loads((cache_dir(pdf) / "manifest.json").read_text())["pages"] for im in row["images"]}
        self.assertEqual({f.name for f in cache_dir(pdf).glob("p*_i*.*")}, named)  # no leftovers
        self.assertTrue(all(name.endswith(".webp") for name in named))  # figures kept as WebP

    def test_real_pptx_upload_extracts_slides_notes_images_and_preserves_download(self):
        path = ROOT / 'curated_content/Data Structures & Algorithms/Lecture 2 - 2024 (1).pptx'
        if not path.exists(): self.skipTest('Local PowerPoint fixture unavailable')
        # Add a notes sentinel to an in-memory copy; the curated original is never changed.
        from io import BytesIO
        from pptx import Presentation
        deck = Presentation(path)
        deck.slides[0].notes_slide.notes_text_frame.text = 'Upload regression: preserve lecturer speaker notes.'
        buffer = BytesIO(); deck.save(buffer)
        data = buffer.getvalue()
        file, = self.reserve([path.name], data)
        uploaded = self.upload(file, data)
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()['page_count'], 22)
        source_id = uploaded.json()['source_id']
        listed = client.get(self.url + '/sources', headers=self.h).json()['sources']
        source = next(s for s in listed if s.get('source_id') == source_id)
        self.assertEqual(source['source_type'], 'pptx')
        with SessionLocal() as db:
            row = db.query(FolderSource).filter_by(source_id=source_id).one()
            pages = load_pages(row.file_path)
        self.assertEqual(len(pages), 22)
        self.assertTrue(any('Speaker notes:' in p['text'] for p in pages))
        self.assertTrue(any('preserve lecturer speaker notes' in p['text'] for p in pages))
        self.assertGreater(sum(len(p['images']) for p in pages), 0)
        downloaded = client.get(self.url + '/sources/' + source_id + '/file', headers=self.h)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, data)
        self.assertIn('presentationml', downloaded.headers['content-type'])


if __name__ == '__main__': unittest.main(verbosity=2)
