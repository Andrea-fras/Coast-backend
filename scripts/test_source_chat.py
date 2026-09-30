#!/usr/bin/env python3
"""Exercise Ask sources through real authenticated HTTP, with isolated data/providers."""
import json
import sys
import uuid
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_http_integrity as fixture
import source_search, source_chat
from database import SessionLocal, FolderSource, SourceSearchIndex, SourceChatTurn, ChatMessage, LearningJob, TutorMemo, StudyFolder
from coast_content_oma.normalized_source import save_pages
import numpy as np
import fitz

client = fixture.client

class SourceChatTests(unittest.TestCase):
    def setUp(self):
        fixture.HttpIntegrity.setUp(self)
        self.scheduler = patch.object(source_search, 'schedule')
        self.scheduler.start(); self.addCleanup(self.scheduler.stop)
        path = Path(fixture.TMP.name) / 'lecture.pdf'
        with fitz.open() as doc:
            for txt in ['Mechanics course. Welcome to the course.', 'Newton: Force equals mass times acceleration. F = ma.', 'The kinetic energy formula is E = 1/2 mv squared.']:
                doc.new_page().insert_text((60,80),txt)
            doc.save(path)
        save_pages(path, [{'page_number':1,'text':'Mechanics introduction'},
            {'page_number':2,'text':'Newton states force equals mass times acceleration. F = ma. A 2 kg body accelerating at 3 m/s squared experiences 6 N.'},
            {'page_number':3,'text':'Kinetic energy equals one half mass times velocity squared. E = 1/2 mv squared.'}])
        with SessionLocal() as db:
            db.add_all([StudyFolder(user_id=1,name='Physics'), FolderSource(user_id=1,folder_name='Physics',source_id='physics',title='Mechanics',filename='lecture.pdf',
                source_type='pdf',page_count=3,raw_text='legacy flat text',file_path=str(path),oma_ingest_status='PENDING'),
                FolderSource(user_id=2,folder_name='Physics',source_id='private',title='Secret',filename='secret.pdf',source_type='pdf',page_count=3,file_path=str(path),raw_text='other student')])
            db.commit()
        self.url='/api/folders/Physics/ask-sources'

    def ask(self, question='What does Newton say about force?', **kwargs):
        body={'message':question,'request_id':str(uuid.uuid4()),**kwargs}
        with patch.object(source_chat,'model_tokens',return_value=iter(['Force equals mass times acceleration. ', '[[S1]]'])), \
             patch('oma_provider.get_student_profile_block',side_effect=AssertionError('OMA must not run')):
            res=client.post(self.url,headers=self.h,json=body)
        events=[json.loads(line[6:]) for line in res.text.splitlines() if line.startswith('data: ')]
        return res, events, body

    def test_ready_before_oma_or_roadmap(self):
        res=client.get(self.url+'/status',headers=self.h)
        self.assertEqual(res.status_code,200,res.text)
        self.assertEqual(res.json()['ready_sources'],1)
        self.assertFalse(res.json()['semantic_ready'])
        results,_=source_search.retrieve(1,'Physics','Newton acceleration force')
        self.assertEqual(results[0]['source_id'],'physics')
        self.assertEqual(results[0]['page'],2)
        self.assertNotIn('private',{r['source_id'] for r in results})

    def test_stream_persists_citations_without_mastery_and_retry_is_idempotent(self):
        res,events,body=self.ask()
        self.assertEqual(res.status_code,200,res.text)
        final=events[-1];self.assertTrue(final['done'],events)
        self.assertEqual(final['citations'][0]['page'],2)
        cid=final['conversation_id']
        history=client.get(self.url+'/history',headers=self.h,params={'conversation_id':cid}).json()
        self.assertEqual(len(history),2)
        self.assertEqual(history[1]['citations'][0]['source_id'],'physics')
        with patch.object(source_chat,'model_tokens',side_effect=AssertionError('cached retry should not call model')):
            retry=client.post(self.url,headers=self.h,json=body)
        self.assertIn('"done": true',retry.text)
        with SessionLocal() as db:
            self.assertEqual(db.query(ChatMessage).count(),2)
            self.assertEqual(db.query(LearningJob).count(),0)
            self.assertEqual(db.query(TutorMemo).count(),0)
        self.assertEqual(len(client.get(self.url+'/conversations',headers=self.h).json()),1)

    def test_ownership_folder_scope_and_invalid_request_identity(self):
        _,events,body=self.ask();cid=events[-1]['conversation_id']
        self.assertEqual(client.get(self.url+'/history',headers=self.other,params={'conversation_id':cid}).status_code,404)
        self.assertEqual(client.get('/api/folders/Other/ask-sources/history',headers=self.h,params={'conversation_id':cid}).status_code,404)
        other={**body,'conversation_id':cid,'request_id':str(uuid.uuid4())}
        self.assertEqual(client.post(self.url,headers=self.other,json=other).status_code,404)
        self.assertEqual(client.post(self.url,headers=self.h,json={**body,'message':'Changed question'}).status_code,409)
        self.assertEqual(client.post(self.url,json=body).status_code,401)

    def test_pdf_page_is_exact_and_authenticated(self):
        base='/api/folders/Physics/sources/physics/pages/'
        res=client.get(base+'2',headers=self.h)
        self.assertEqual(res.status_code,200)
        self.assertTrue(res.content.startswith(b'\x89PNG'))
        self.assertNotEqual(res.content,client.get(base+'1',headers=self.h).content)
        self.assertEqual(client.get(base+'2',headers=self.other).status_code,404)
        self.assertEqual(client.get(base+'0',headers=self.h).status_code,404)
        self.assertEqual(client.get(base+'4',headers=self.h).status_code,404)

    def test_hybrid_semantic_search_finds_paraphrase(self):
        source_search.status(1,'Physics')
        with SessionLocal() as db:
            row=db.get(SourceSearchIndex,'physics')
            row.vector_count=3;row.dimensions=2
            row.vectors=np.asarray([[0.,1.],[1.,0.],[0.,1.]],dtype='<f4').tobytes();db.commit()
        with patch.dict('os.environ',{'OPENAI_API_KEY':'test'}),patch.object(source_search,'embed',return_value=np.asarray([[1.,0.]],dtype='<f4')):
            passages,coverage=source_search.retrieve(1,'Physics','What determines the push needed to change motion?')
        self.assertEqual(passages[0]['page'],2)
        self.assertEqual(coverage['retrieval'],'hybrid')

    def test_no_match_does_not_invent_answer_or_call_model(self):
        with patch.object(source_chat,'model_tokens',side_effect=AssertionError('No evidence')):
            claim=source_chat.begin(1,'Physics','quasars superconductivity',str(uuid.uuid4()))
            events=list(source_chat.stream_answer(1,'Physics','quasars superconductivity',claim))
        self.assertTrue(events[-1]['done'])
        self.assertEqual(events[-1]['citations'],[])
        self.assertIn('couldn’t find',events[-1]['reply'])

    def test_invalid_citation_rejected(self):
        claim=source_chat.begin(1,'Physics','Newton force',str(uuid.uuid4()))
        with patch.object(source_chat,'model_tokens',return_value=iter(['Invented answer [[S999]]'])):
            result=list(source_chat.stream_answer(1,'Physics','Newton force',claim))[-1]
        self.assertEqual(result['citations'],[])
        self.assertNotIn('Invented answer',result['reply'])

    def test_grouped_citations_survive_stream_save_history_and_retry(self):
        evidence = [dict(id=sid, source_id='physics', title='Mechanics', filename='lecture.pdf',
                         source_type='pdf', page=page, page_count=3, text='Physics evidence')
                    for sid, page in [('S6',2),('S7',3)]]
        coverage = dict(ready_sources=1,total_sources=1,semantic_ready=True)
        body = dict(message='Compare these ideas',request_id=str(uuid.uuid4()))
        with patch.object(source_search,'retrieve',return_value=(evidence,coverage)), \
             patch.object(source_chat,'model_tokens',return_value=iter(['Comparison **[[S6], ', '[S7]]**. Extra [[S999]].'])):
            result=client.post(self.url,headers=self.h,json=body)
        events=[json.loads(line[6:]) for line in result.text.splitlines() if line.startswith('data: ')]
        final=events[-1]
        self.assertTrue(final['done'])
        self.assertIn('**[[S6]] [[S7]]**',final['reply'])
        self.assertNotIn('S999',final['reply'])
        self.assertEqual([(c['id'],c['page']) for c in final['citations']],[('S6',2),('S7',3)])
        history=client.get(self.url+'/history',headers=self.h,params={'conversation_id':final['conversation_id']}).json()
        self.assertEqual(history[-1]['content'],final['reply'])
        self.assertEqual([c['id'] for c in history[-1]['citations']],['S6','S7'])
        with patch.object(source_chat,'model_tokens',side_effect=AssertionError('Do not regenerate a saved answer')):
            retry=client.post(self.url,headers=self.h,json=body)
        replay=[json.loads(line[6:]) for line in retry.text.splitlines() if line.startswith('data: ')][-1]
        self.assertEqual(replay['citations'],final['citations'])

    def test_provider_failure_retries_same_question_and_remount_has_retry(self):
        rid=str(uuid.uuid4());claim=source_chat.begin(1,'Physics','Newton force',rid)
        with patch.object(source_chat,'model_tokens',side_effect=RuntimeError('provider private details')):
            events=list(source_chat.stream_answer(1,'Physics','Newton force',claim))
        self.assertIn('error',events[-1]);self.assertNotIn('private details',events[-1]['error'])
        hist=source_chat.history(1,'Physics',claim['conversation_id'])
        self.assertEqual(hist[-1]['status'],'failed')
        _,events,_=self.ask('Newton force',request_id=rid,conversation_id=claim['conversation_id'])
        self.assertTrue(events[-1]['done'])
        with SessionLocal() as db: self.assertEqual(db.query(ChatMessage).count(),2)

    def test_source_removal_invalidates_retrieval_and_marks_saved_citation(self):
        _,events,_=self.ask();cid=events[-1]['conversation_id']
        res=client.delete('/api/folders/Physics/sources/physics',headers=self.h)
        self.assertEqual(res.status_code,200,res.text)
        results,_=source_search.retrieve(1,'Physics','Newton force')
        self.assertEqual(results,[])
        history=source_chat.history(1,'Physics',cid)
        self.assertFalse(history[-1]['citations'][0]['available'])
        with SessionLocal() as db: self.assertIsNone(db.get(SourceSearchIndex,'physics'))

    def test_busy_question_and_cross_mode_conversation_rejected(self):
        claim=source_chat.begin(1,'Physics','Newton force',str(uuid.uuid4()))
        with self.assertRaises(Exception) as err: source_chat.begin(1,'Physics','More',str(uuid.uuid4()),claim['conversation_id'])
        self.assertEqual(err.exception.status_code,409)
        with SessionLocal() as db:
            db.add(ChatMessage(user_id=1,context_type='lesson',context_id='Physics',conversation_id='lesson-one',role='user',content='Explain'));db.commit()
        with self.assertRaises(Exception) as err: source_chat.begin(1,'Physics','Force',str(uuid.uuid4()),'lesson-one')
        self.assertEqual(err.exception.status_code,404)

    def test_followup_retains_previous_evidence(self):
        evidence,_=source_search.retrieve(1,'Physics','Give an example of that', 'Newton acceleration', [('physics',2)])
        self.assertEqual(evidence[0]['page'],2)

    def test_installed_gemini_client_without_close_streams_successfully(self):
        from types import SimpleNamespace
        fake = SimpleNamespace(models=SimpleNamespace(generate_content_stream=lambda **kwargs: iter([SimpleNamespace(text='Supported answer [[S1]]')])))
        with patch('google.genai.Client',return_value=fake),patch('tutor.CHAT_PROVIDER','gemini'):
            self.assertEqual(''.join(source_chat.model_tokens([{'role':'system','content':'Sources'},{'role':'user','content':'Question'}])), 'Supported answer [[S1]]')

    def test_old_attempt_cannot_overwrite_a_retried_question(self):
        rid=str(uuid.uuid4());old=source_chat.begin(1,'Physics','Newton force',rid)
        with SessionLocal() as db:
            run=db.get(SourceChatTurn,(1,rid));run.started_at=0;db.commit()
        fresh=source_chat.begin(1,'Physics','Newton force',rid)
        with patch.object(source_chat,'model_tokens',return_value=iter(['Stale [[S1]]'])):
            events=list(source_chat.stream_answer(1,'Physics','Newton force',old))
        self.assertIn('error',events[-1])
        with SessionLocal() as db:
            run=db.get(SourceChatTurn,(1,rid));self.assertEqual(run.status,'running');self.assertEqual(run.started_at,fresh['attempt_at'])

    def test_real_pptx_upload_is_immediately_searchable(self):
        from pptx import Presentation
        from io import BytesIO
        deck=Presentation()
        for title,body in [('Lecture overview','Basic physics'),('Momentum','Momentum equals mass multiplied by velocity. It is conserved in an isolated system.')]:
            slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body
        buf=BytesIO();deck.save(buf);data=buf.getvalue();upload_id=str(uuid.uuid4())
        reg=client.post('/api/folders/Physics/uploads',headers=self.h,json={'files':[{'upload_id':upload_id,'filename':'Momentum.pptx','size_bytes':len(data)}]})
        self.assertEqual(reg.status_code,200,reg.text)
        with patch('oma_provider.is_oma_enabled',return_value=True),patch('learning_jobs.wake'):
            uploaded=client.post('/api/folders/Physics/upload',headers=self.h,data={'upload_id':upload_id},files={'file':('Momentum.pptx',data,'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        evidence,_=source_search.retrieve(1,'Physics','How is momentum defined?')
        self.assertEqual(evidence[0]['source_id'],uploaded.json()['source_id'])
        self.assertEqual(evidence[0]['page'],2)
        page=client.get('/api/folders/Physics/sources/'+uploaded.json()['source_id']+'/pages/2',headers=self.h)
        self.assertIn('Momentum equals',page.json()['text'])

    def test_pending_files_are_in_answer_coverage(self):
        from database import SourceUpload
        with SessionLocal() as db:
            db.add(SourceUpload(user_id=1,folder_name='Physics',upload_id='pending',filename='next.pdf',size_bytes=100,status='queued',expires_at=9999999999));db.commit()
        _,events,_=self.ask()
        self.assertEqual(events[-1]['coverage']['pending_uploads'],1)
        self.assertEqual(events[-1]['coverage']['ready_sources'],1)

    def test_powerpoint_preview_uses_slide_numbers(self):
        with SessionLocal() as db:
            row=db.query(FolderSource).filter_by(source_id='physics').first();row.source_type='pptx';db.commit()
        res=client.get('/api/folders/Physics/sources/physics/pages/2',headers=self.h)
        self.assertEqual(res.status_code,200,res.text)
        self.assertEqual(res.json()['page'],2)
        self.assertIn('Newton',res.json()['text'])

if __name__=='__main__':unittest.main()
