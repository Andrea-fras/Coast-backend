#!/usr/bin/env python3
"""Offline launch regressions. Every database lives in a temporary directory."""
import os
import sys
import tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TEMP = tempfile.TemporaryDirectory(prefix="coast-regression-")
os.environ['DATABASE_PATH'] = str(Path(TEMP.name) / 'app.db')
os.environ['OMA_DB_PATH'] = str(Path(TEMP.name) / 'oma.db')
for key in ('OPENAI_API_KEY','GEMINI_API_KEY','ANTHROPIC_API_KEY'):
    os.environ.pop(key, None)
import json
import time
import unittest
from unittest.mock import patch
import dotenv
dotenv.load_dotenv = lambda *a, **k: None  # Never load developer credentials in offline tests.
from database import Base, engine, SessionLocal, User, CourseOutline, LearningJob, PlacementTestSession
import placement
import learning_jobs
from coast_content_oma.stores.db import connect_db, transaction
from coast_content_oma.stores.base import MemoryItem
from coast_content_oma.student.stores.episode import EpisodeStore
from coast_content_oma.stores.concept_alias import ConceptAliasStore
from coast_content_oma.student.prompt_budget import course_profile
from notes_service import sanitize_notes, notes_revision

class LaunchIntegrity(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(engine)
        Base.metadata.create_all(engine)
        with SessionLocal() as db:
            db.add(User(id=1, email='fixture@example.invalid', name='Fixture'))
            db.add(User(id=2, email='other@example.invalid', name='Other'))
            db.add(CourseOutline(user_id=1, folder_name='Physics', current_section=0,
                total_sections=3, outline_json=json.dumps([{'title': 'A'},{'title':'B'},{'title':'C'}])))
            db.commit()

    def _placement(self, turns):
        """Run a placement from section 1 towards section 3, one (student, Pedro) turn at a time."""
        import placement
        from database import ChatMessage
        placement.begin(1, 'Physics', 2, 'placement-fixture')
        state = None
        for student, reply in turns:
            with SessionLocal() as db:
                db.add(ChatMessage(user_id=1, conversation_id='placement-fixture', role='user', content=student,
                                   context_type='test_out', context_id='Physics'))
                db.add(ChatMessage(user_id=1, conversation_id='placement-fixture', role='pedro', content=reply,
                                   context_type='test_out', context_id='Physics'))
                db.commit()
            state = placement.record_turn(1, 'placement-fixture', reply)
        return state

    def test_placement_pass_counts_when_pedro_forgets_the_grade_tag(self):
        state = self._placement([('Check what I know', 'First question about A?'),
                                 ('my answer', 'Right. [PLACEMENT_PASSED: 1] Now B?'),
                                 ('my answer', 'Right. [ANSWER_CORRECT: b] [PLACEMENT_PASSED: 2]')])
        self.assertEqual((state['passed_count'], state['done'], state['can_apply']), (2, True, True))

    def test_placement_pass_needs_an_answer_and_counts_once_per_reply(self):
        state = self._placement([('Check what I know', 'You know A. [PLACEMENT_PASSED: 1]')])
        self.assertEqual(state['passed_count'], 0)  # Pedro's opening message can't pass anything
        state = self._placement([('my answer', 'Right. [PLACEMENT_PASSED: 1] [PLACEMENT_PASSED: 2]')])
        self.assertEqual(state['passed_count'], 1)  # one answer, one section

    def test_placement_pass_ignored_beside_a_wrong_or_helped_answer(self):
        state = self._placement([('Check what I know', 'First question about A?'),
                                 ('my answer', 'Not quite. [ANSWER_WRONG: a] [PLACEMENT_PASSED: 1]'),
                                 ('my answer', 'With my hint, yes. [ANSWER_CORRECT: a | hinted] [PLACEMENT_PASSED: 1]')])
        self.assertEqual(state['passed_count'], 0)

    def test_completion_rollback_is_atomic(self):
        with SessionLocal() as db:
            learning_jobs.enqueue(db,1,'Physics',0)
            db.rollback()
        with SessionLocal() as db:
            self.assertEqual(db.query(LearningJob).count(),0)

    def test_overlapping_sections_both_run_and_duplicate_is_ignored(self):
        with SessionLocal() as db:
            first=learning_jobs.enqueue(db,1,'Physics',0)
            self.assertEqual(first,learning_jobs.enqueue(db,1,'Physics',0))
            learning_jobs.enqueue(db,1,'Physics',1)
            db.commit()
        seen=[]
        while learning_jobs.run_one(lambda job,p: seen.append(p['section_index'])):
            pass
        self.assertEqual(sorted(seen),[0,1])
        with SessionLocal() as db:
            self.assertEqual(db.query(LearningJob).filter_by(status='done').count(),2)

    def test_expired_worker_lease_recovered_and_old_worker_cannot_finish(self):
        with SessionLocal() as db:
            learning_jobs.enqueue(db,1,'Physics',0)
            db.commit()
        job_id,old_token,_=learning_jobs.claim()
        with SessionLocal() as db:
            db.get(LearningJob,job_id).lease_until=time.time()-1
            db.commit()
        _,new_token,_=learning_jobs.claim()
        self.assertNotEqual(old_token,new_token)
        learning_jobs.finish(job_id,old_token)
        with SessionLocal() as db:
            self.assertEqual(db.get(LearningJob,job_id).status,'running')
        learning_jobs.finish(job_id,new_token)

    def test_indexing_takes_turns_between_students(self):
        # Student 1 uploaded three files and one is already being indexed; student 2's single
        # file is just as urgent, so it goes next instead of waiting behind student 1's upload.
        now = time.time()
        with SessionLocal() as db:
            for n, (uid, status) in enumerate([(1, 'running'), (1, 'queued'), (1, 'queued'), (2, 'queued')]):
                db.add(LearningJob(id=f'src-{n}', status=status, attempts=0, available_at=0,
                                   lease_until=now + 60 if status == 'running' else 0,
                                   payload_json=json.dumps({'kind': 'source_ingest', 'user_id': uid, 'folder': 'Physics',
                                                            'source_id': f's{n}', 'path': f'/x/s{n}.pdf'})))
            db.commit()
        claimed = learning_jobs.claim('source')
        self.assertEqual(claimed[2]['user_id'], 2)
        self.assertEqual(learning_jobs.claim('source')[2]['source_id'], 's1')  # then student 1's next file

    def test_failed_job_keeps_payload_and_retries(self):
        with SessionLocal() as db:
            job_id=learning_jobs.enqueue(db,1,'Physics',2)
            db.commit()
        with patch.object(learning_jobs.log,'exception'):
            learning_jobs.run_one(lambda *_: (_ for _ in ()).throw(RuntimeError('fixture failure')))
        with SessionLocal() as db:
            row=db.get(LearningJob,job_id)
            self.assertEqual(row.status,'queued')
            self.assertEqual(json.loads(row.payload_json)['section_index'],2)
            self.assertIn('fixture failure',row.last_error)

    def test_placement_requires_real_pass_and_matching_outline_and_student(self):
        placement.begin(1,'Physics',2,'conv_fixture')
        with SessionLocal() as db:
            outline=db.query(CourseOutline).first()
            with self.assertRaises(ValueError):
                placement.authorize(db,1,'Physics',2,'conv_fixture',outline)
        self.assertTrue(placement.record_pass(1,'conv_fixture'))
        with SessionLocal() as db:
            outline=db.query(CourseOutline).first()
            self.assertTrue(placement.authorize(db,1,'Physics',2,'conv_fixture',outline).passed)
            for uid,target in [(2,2),(1,1)]:
                with self.assertRaises(ValueError):
                    placement.authorize(db,uid,'Physics',target,'conv_fixture',outline)
            outline.outline_json='[{"title":"Changed"}]'
            with self.assertRaises(ValueError):
                placement.authorize(db,1,'Physics',2,'conv_fixture',outline)

    def test_fts_update_has_one_current_posting(self):
        with tempfile.TemporaryDirectory() as d:
            store=EpisodeStore(Path(d)/'oma.db')
            item=MemoryItem(id='test',namespace='u1',store='episode',content='oldword')
            store._insert(item)
            item.content='newword'
            store._insert(item)
            with connect_db(store.db_path) as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM episode_items_fts').fetchone()[0],1)
                self.assertEqual(conn.execute("SELECT count(*) FROM episode_items_fts WHERE episode_items_fts MATCH 'oldword'").fetchone()[0],0)

    def test_nested_store_schema_does_not_commit_partial_projection(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'oma.db'
            store=EpisodeStore(path)
            with self.assertRaises(RuntimeError):
                with transaction(path):
                    store._insert(MemoryItem(id='rollback',namespace='u1',store='episode',content='test'))
                    ConceptAliasStore(path)
                    raise RuntimeError('crash after partial write')
            self.assertIsNone(store.get('rollback'))

    def test_projection_crash_rolls_back_and_retry_records_once(self):
        from contextlib import ExitStack
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as d, ExitStack() as patches:
            store = EpisodeStore(Path(d) / 'projection.db')
            recorder = SimpleNamespace(record_episode=lambda *a, **kw: store.write('u1', 'completion'))
            patches.enter_context(patch('map_world.claim_section_reward'))
            patches.enter_context(patch('oma_provider.is_student_enabled', return_value=True))
            patches.enter_context(patch('oma_provider._student_orchestrator', return_value=SimpleNamespace(episodes=store)))
            patches.enter_context(patch('oma_provider._student_recorder_singleton', return_value=recorder))
            patches.enter_context(patch('oma_provider._run_course_consolidation'))
            patches.enter_context(patch('evaluator.fetch_section_transcript', return_value=([], 'Student independently explains work.')))
            patches.enter_context(patch('lesson.get_section_concept_refs', return_value=[]))
            patches.enter_context(patch('evaluator.evaluate_section_transcript', return_value={}))
            def interrupted_apply(*args):
                store.write('u1', 'partial evaluation')
                raise RuntimeError('simulated crash during projection')
            apply = patches.enter_context(patch('evaluator.apply_evaluation', side_effect=interrupted_apply))
            with SessionLocal() as db:
                event = learning_jobs.enqueue(db, 1, 'Physics', 0)
                db.commit()
                payload = json.loads(db.get(LearningJob, event).payload_json)
            payload['through_message_id'] = 1
            with self.assertRaises(RuntimeError):
                learning_jobs.project(event, payload)
            with connect_db(store.db_path) as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM episode_items').fetchone()[0], 0)
                self.assertEqual(conn.execute('SELECT count(*) FROM applied_learning_events').fetchone()[0], 0)
            apply.side_effect = lambda *a: store.write('u1', 'complete evaluation')
            learning_jobs.project(event, payload)
            learning_jobs.project(event, payload)
            with connect_db(store.db_path) as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM episode_items').fetchone()[0], 2)
                self.assertEqual(conn.execute('SELECT count(*) FROM applied_learning_events').fetchone()[0], 1)
            self.assertEqual(apply.call_count, 2)  # failed attempt plus one successful retry

    def test_student_writes_do_not_request_unused_embeddings(self):
        with tempfile.TemporaryDirectory() as d, patch.dict(os.environ, {'OPENAI_API_KEY': 'fixture-only'}):
            store = EpisodeStore(Path(d) / 'oma.db')
            with patch('coast_content_oma.stores._semantic_base.provider_capacity.call') as api:
                store.write('u1', 'one memory')
                store.write_items_bulk([MemoryItem(id='bulk', namespace='u1', store='episode', content='another memory')])
                api.assert_not_called()

    def test_relevant_memory_survives_long_history(self):
        profile={'focused_mastery':[{'concept_id':'orm','name':'ORM','score':0.8,'successes':3}],
            'golden_moments':[{'concept_ids':['orm'],'text':'Translator analogy worked.'}],
            'identity_traits':[{'text':'Prefers concrete examples.'}],
            'progress_ledger':{'completed_sections':[{'title': 'Long completed section '+str(i)+' x'*200} for i in range(100)]}}
        block=course_profile(profile)
        self.assertLessEqual(len(block),1200)
        for text in ('never invent','Current concept: ORM','Translator analogy','Prefers concrete'):
            self.assertIn(text,block)
        profile['golden_moments'][0]['concept_ids']=['unrelated']
        self.assertNotIn('Translator analogy',course_profile(profile))

    def test_bulk_reindex_replaces_old_fts_content(self):
        with tempfile.TemporaryDirectory() as d:
            store=EpisodeStore(Path(d)/'oma.db')
            items=[MemoryItem(id=str(i),namespace='u1',store='episode',content='oldword') for i in range(3)]
            store.write_items_bulk(items)
            for item in items: item.content='newword'
            store.write_items_bulk(items)
            with connect_db(store.db_path) as conn:
                self.assertEqual(conn.execute('SELECT count(*) FROM episode_items_fts').fetchone()[0],3)
                self.assertEqual(conn.execute("SELECT count(*) FROM episode_items_fts WHERE episode_items_fts MATCH 'oldword'").fetchone()[0],0)

    def test_namespace_migration_preserves_legacy_and_rename(self):
        from database import CourseIdentity
        from coast_content_oma.course_identity import initialize,namespace_key,new_key,display_name
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'oma.db'
            store=EpisodeStore(path)
            store.write('u1__student__physics','legacy memory')
            initialize(path)
            self.assertEqual(namespace_key(1,'Physics'),'physics')
            with SessionLocal() as db:
                db.get(CourseIdentity,(1,'Physics')).folder_name='Physics renamed'
                db.commit()
            self.assertEqual(namespace_key(1,'Physics renamed'),'physics')
            self.assertEqual(display_name(1,'physics'),'Physics renamed')
            self.assertNotEqual(new_key('A-B'),new_key('A B'))
            self.assertNotEqual(new_key('物理'),new_key('化学'))

    def test_legacy_namespace_collision_does_not_assign_wrong_student_history(self):
        from database import CourseIdentity,StudyFolder
        from coast_content_oma.course_identity import initialize
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'oma.db'
            EpisodeStore(path).write('u1__student__a_b','ambiguous old memory')
            with SessionLocal() as db:
                db.add_all([StudyFolder(user_id=1,name='A-B'),StudyFolder(user_id=1,name='A B')])
                db.commit()
            with self.assertRaises(RuntimeError): initialize(path)
            with SessionLocal() as db: self.assertEqual(db.query(CourseIdentity).count(),0)

    def test_cross_course_recall_is_owner_scoped_and_evidence_filtered(self):
        from coast_content_oma.student.recall import recall_memories
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'oma.db';store=EpisodeStore(path)
            for uid,kind,text in [(1,'section_completed','Translator analogy in databases'),
                                   (2,'section_completed','Translator private other student'),
                                   (1,'chat_turn','Translator unverified inference')]:
                store.write(f'u{uid}__student__old_course',text,store_specific={'episode_type':kind})
            recalled=recall_memories(path,1,'Remember our translator analogy?')
            self.assertEqual(len(recalled),1)
            self.assertIn('databases',recalled[0]['text'])

    def test_normalized_source_cache_invalidates_when_bytes_change(self):
        from coast_content_oma.normalized_source import save_pages,load_pages
        from PIL import Image
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'source.pdf';path.write_bytes(b'original fixture')
            save_pages(path,[{'page_number':1,'text':'A diagram','images':[{'idx':0,'pil_image':Image.new('RGB',(100,100))}]}])
            pages=load_pages(path)
            self.assertEqual(pages[0]['text'],'A diagram')
            self.assertEqual(pages[0]['images'][0]['pil_image'].size,(100,100))
            self.assertEqual(load_pages(path,False)[0]['images'],[])
            path.write_bytes(b'replaced fixture')
            self.assertIsNone(load_pages(path))

    def test_rolling_summary_reuses_digest_without_losing_new_turns(self):
        from database import ChatMessage
        from conversation_memory import context
        def append(count):
            with SessionLocal() as db:
                for i in range(count): db.add(ChatMessage(user_id=1,conversation_id='c',role='user',content=str(i),context_type='lesson'))
                db.commit()
        append(14)
        summarize=unittest.mock.Mock(return_value='Earlier evidence')
        summary,recent=context(1,'c',summarize)
        self.assertEqual(len(recent),6)
        self.assertEqual(len(summarize.call_args.args[0]),8)
        append(2)
        summary,recent=context(1,'c',summarize)
        self.assertEqual(summarize.call_count,1,'Do not pay to summarize unchanged old turns again')
        self.assertEqual(len(recent),8,'New unsummarized messages must remain in context')
        append(6)
        summary,recent=context(1,'c',summarize)
        self.assertEqual(summarize.call_count,2)
        self.assertEqual(summarize.call_args.args[1],'Earlier evidence')
        self.assertEqual(context(2,'c',summarize),(None,[]))

    def test_source_deletion_clears_live_indexes_but_preserves_student_evidence(self):
        from coast_content_oma.stores.content import ContentStore
        from coast_content_oma.source_lifecycle import remove_source_material
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'oma.db';store=ContentStore(path);episodes=EpisodeStore(path)
            gone=store.write('u1__physics','Deleted source definition',source_doc_id='doc_src_fixture')
            kept=store.write('u1__physics','Other source',source_doc_id='doc_src_other')
            other=store.write('u2__physics','Another student',source_doc_id='doc_src_fixture')
            episode=episodes.write('u1__student__physics','Historical completion')
            remove_source_material(path,'u1__physics','src_fixture')
            self.assertIsNone(store.get(gone.id));self.assertIsNotNone(store.get(kept.id));self.assertIsNotNone(store.get(other.id))
            self.assertIsNotNone(episodes.get(episode.id))
            with connect_db(path) as conn:
                self.assertEqual(conn.execute("SELECT count(*) FROM content_items_fts WHERE content_items_fts MATCH 'Deleted'").fetchone()[0],0)

    def test_provider_capacity_reserves_live_slot_and_releases_failed_stream(self):
        import provider_capacity as capacity
        gate=capacity.Capacity(concurrent=2,background=1)
        with gate.slot('background'):
            with self.assertRaises(TimeoutError):
                with gate.slot('background',timeout=0): pass
            with gate.slot('interactive',timeout=0): self.assertEqual(gate.active,2)
        self.assertEqual(gate.active,0)
        with patch.object(capacity,'gate',return_value=gate):
            def failing():
                yield 'first token'
                raise RuntimeError('connection lost')
            stream=capacity.stream('gemini',failing)
            self.assertEqual(next(stream),'first token')
            with self.assertRaises(RuntimeError): next(stream)
            self.assertEqual(gate.active,0)
            stream=capacity.stream('gemini',lambda:iter(['a','b']))
            next(stream);stream.close()
            self.assertEqual(gate.active,0)

    def test_evaluator_window_excludes_old_attempt_and_retains_final_answer(self):
        import evaluator
        from database import ChatMessage
        with SessionLocal() as db:
            old=ChatMessage(user_id=1,conversation_id='old',role='user',content='OLD_MISCONCEPTION',context_type='lesson',context_id='Physics',section_index=0)
            db.add(old);db.flush();cutoff=old.id
            for role,content in [('pedro','Long explanation '*500),('user','FINAL_CORRECT_REASONING')]:
                db.add(ChatMessage(user_id=1,conversation_id='new',role=role,content=content,context_type='lesson',context_id='Physics',section_index=0))
            db.commit()
        _,transcript=evaluator.fetch_section_transcript(1,'Physics',0,after_message_id=cutoff,max_chars=1000)
        self.assertNotIn('OLD_MISCONCEPTION',transcript)
        self.assertIn('FINAL_CORRECT_REASONING',transcript)
        self.assertIn('omitted',transcript)
        self.assertLessEqual(len(transcript),1000)
        with self.assertRaises(RuntimeError): evaluator.evaluate_section_transcript('Student: answer',0,'A',[],require_model=True)

    def test_pdf_and_pptx_share_page_text_and_image_contract(self):
        import fitz
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        from coast_content_oma.extraction import extract_pages
        with tempfile.TemporaryDirectory() as d:
            path=Path(d);image=path/'image.png';Image.new('RGB',(120,120),'blue').save(image)
            pdf=fitz.open();page=pdf.new_page();page.insert_text((30,30),'Energy transfer');page.insert_image(fitz.Rect(30,70,150,190),filename=str(image));pdf.save(path/'lecture.pdf');pdf.close()
            presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1)).text='Energy transfer'
            slide.shapes.add_picture(str(image),Inches(1),Inches(3))
            presentation.save(path/'lecture.pptx')
            for name in ('lecture.pdf','lecture.pptx'):
                pages=extract_pages(path/name,use_cache=False)
                self.assertEqual(pages[0]['page_number'],1)
                self.assertIn('Energy transfer',pages[0]['text'])
                self.assertEqual(len(pages[0]['images']),1)

    def test_note_html_cannot_store_executable_attributes_or_urls(self):
        html=sanitize_notes('<p onclick="x()">Hello</p><img src="x" onerror="alert(1)"><a href="javascript:alert(1)">link</a><script>alert(1)</script>')
        for value in ('onclick','onerror','javascript:','<script'):
            self.assertNotIn(value,html)
        self.assertIn('<p>Hello</p>',html)
        self.assertNotEqual(notes_revision('a'),notes_revision('b'))

    def test_image_token_cannot_be_used_as_account_token(self):
        from auth import create_image_token,decode_image_token,decode_access_token
        token=create_image_token(1)
        self.assertEqual(decode_image_token(token)['sub'],'1')
        self.assertIsNone(decode_access_token(token))

if __name__=='__main__':
    unittest.main(verbosity=2)
